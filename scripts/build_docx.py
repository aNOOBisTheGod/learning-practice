import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Mm, Pt, RGBColor

root = Path(__file__).resolve().parent.parent
doc = Document()
sec = doc.sections[0]
sec.page_width, sec.page_height = Mm(210), Mm(297)
sec.left_margin, sec.right_margin = Mm(30), Mm(10)
sec.top_margin = sec.bottom_margin = Mm(20)
sec.footer_distance = Mm(10)
sec.different_first_page_header_footer = True
for name in ['Normal', 'Title', 'Heading 1', 'Heading 2', 'TOC Heading']:
    style = doc.styles[name]
    style.font.name = 'Times New Roman'
    style.font.size = Pt(14)
    style.font.color.rgb = RGBColor(0, 0, 0)
    rpr = style.element.get_or_add_rPr()
    fonts = rpr.find(qn('w:rFonts'))
    for key in list(fonts.attrib):
        if 'theme' in key.lower(): del fonts.attrib[key]
    for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
        fonts.set(qn('w:' + key), 'Times New Roman')
    for el in list(style.element.xpath('./w:pPr/w:pBdr')):
        el.getparent().remove(el)
    fmt = style.paragraph_format
    fmt.line_spacing = 1.5
    fmt.space_after = Pt(0)
    fmt.space_before = Pt(0)
    fmt.widow_control = True
normal = doc.styles['Normal'].paragraph_format
normal.first_line_indent = Mm(12.5)
normal.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
for name in ['Heading 1', 'Heading 2']:
    fmt = doc.styles[name].paragraph_format
    fmt.first_line_indent = Mm(0)
    fmt.keep_with_next = True
    doc.styles[name].font.bold = True
    fmt.space_after = Pt(12)
doc.styles['Heading 1'].paragraph_format.page_break_before = True
doc.styles['Heading 1'].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER
doc.styles['Heading 2'].paragraph_format.space_before = Pt(10)
lang = OxmlElement('w:lang'); lang.set(qn('w:val'), 'ru-RU')
doc.styles['Normal'].element.get_or_add_rPr().append(lang)

def para(text, align=WD_ALIGN_PARAGRAPH.CENTER, before=0, after=5, bold=False, style=None):
    p = doc.add_paragraph(text, style)
    p.alignment = align
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_before = Pt(before)
    p.paragraph_format.space_after = Pt(after)
    if bold:
        for r in p.runs: r.bold = True
    return p

def field(p, instruction, placeholder=''):
    r = p.add_run()
    begin = OxmlElement('w:fldChar'); begin.set(qn('w:fldCharType'), 'begin')
    instr = OxmlElement('w:instrText'); instr.set(qn('xml:space'), 'preserve'); instr.text = instruction
    sep = OxmlElement('w:fldChar'); sep.set(qn('w:fldCharType'), 'separate')
    text = OxmlElement('w:t'); text.text = placeholder
    end = OxmlElement('w:fldChar'); end.set(qn('w:fldCharType'), 'end')
    for el in [begin, instr, sep, text, end]: r._r.append(el)

p = sec.footer.paragraphs[0]
p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.paragraph_format.first_line_indent = Mm(0)
field(p, ' PAGE ', '2')
for text in ['МИНОБРНАУКИ РОССИИ',
             'Федеральное государственное бюджетное образовательное учреждение',
             'высшего образования',
             '«МИРЭА – Российский технологический университет»',
             'РТУ МИРЭА']:
    p = para(text, after=0)
    p.paragraph_format.line_spacing = 1
para('ОТЧЁТ', before=90, bold=True, style='Title')
para('по учебной практике', after=0)
para('КТ1 — Исследовательский этап', after=25)
para('Тема: «Сборка и тестирование серверного оборудования»', bold=True, after=35)
para('Место практики: ООО «Бюджетные и Финансовые Технологии»', WD_ALIGN_PARAGRAPH.LEFT, after=25)
for text in ['Выполнил: студент группы ИКБО-10-24',
             'Новожилов Василий Алексеевич',
             'Руководитель от университета:',
             'Синицын Анатолий Васильевич',
             'Руководитель от предприятия:',
             'Жиркова М. И.']:
    p = para(text, WD_ALIGN_PARAGRAPH.LEFT, after=0)
    p.paragraph_format.left_indent = Mm(70)
    p.paragraph_format.line_spacing = 1
para('Москва 2026', before=70)
doc.add_page_break()
para('Содержание', bold=True, after=15)
headings = [block[2:] for block in (root / 'docs/report.txt').read_text().strip().split('\n\n') if block.startswith('# ')]
headings.append('Список использованных источников')
for i, heading in enumerate(headings):
    p = doc.add_paragraph(heading + '\t')
    field(p, f' PAGEREF section_{i} \\h ', '')
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.tab_stops.add_tab_stop(Mm(170), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
    p.paragraph_format.space_after = Pt(8)
def section_heading(text):
    p = doc.add_heading(text, 1)
    index = headings.index(text)
    start = OxmlElement('w:bookmarkStart')
    start.set(qn('w:id'), str(index))
    start.set(qn('w:name'), f'section_{index}')
    end = OxmlElement('w:bookmarkEnd'); end.set(qn('w:id'), str(index))
    p._p.insert(1, start); p._p.append(end)
    return p

for block in (root / 'docs/report.txt').read_text().strip().split('\n\n'):
    if block.startswith('# '): section_heading(block[2:])
    elif block.startswith('## '): doc.add_heading(block[3:], 2)
    else: doc.add_paragraph(block.replace('\n', ' '))
section_heading('Список использованных источников')
for s in json.loads((root / 'docs/sources.json').read_text()):
    p = doc.add_paragraph(f"{s['id']}. {s['reference']} - URL: {s['url']} (дата обращения: 05.10.2026).")
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p.paragraph_format.first_line_indent = Mm(0)
    p.paragraph_format.space_after = Pt(10)
doc.core_properties.author = 'Новожилов Василий Алексеевич'
doc.core_properties.title = 'Сборка и тестирование серверного оборудования'
doc.core_properties.subject = 'Отчёт по учебной практике КТ1'
doc.core_properties.comments = ''
out = root / 'submission/Отчет_ИКБО-10-24_НовожиловВА_КТ1.docx'
for p in doc.paragraphs:
    for run in p.runs:
        run.font.name = 'Times New Roman'
        run.font.size = Pt(14)
        run.font.color.rgb = RGBColor(0, 0, 0)
        fonts = run._r.get_or_add_rPr().find(qn('w:rFonts'))
        for key in list(fonts.attrib):
            if 'theme' in key.lower(): del fonts.attrib[key]
        for key in ['ascii', 'hAnsi', 'eastAsia', 'cs']:
            fonts.set(qn('w:' + key), 'Times New Roman')
    for border in list(p._p.xpath('./w:pPr/w:pBdr')):
        border.getparent().remove(border)
update = OxmlElement('w:updateFields')
update.set(qn('w:val'), 'true')
doc.settings.element.append(update)
doc.save(out)
print(out)
